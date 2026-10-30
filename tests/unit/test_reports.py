import io

import pandas as pd
from openpyxl import load_workbook

from reports.exports import (
    backup_database,
    build_workbook,
    get_data_quality_summary,
    get_xirr_audit,
    restore_database,
    workbook_to_bytes,
)


def test_workbook():
    wb = build_workbook({"A": pd.DataFrame({"x": [1,2]}), "B": pd.DataFrame({"y": [3]})}, metadata={"k": "v"})
    b = workbook_to_bytes(wb)
    wb2 = load_workbook(io.BytesIO(b))
    assert "Metadata" in wb2.sheetnames
    assert "A" in wb2.sheetnames and "B" in wb2.sheetnames

def test_data_quality():
    txs = [{"trade_date": "2024-01-01", "isin": "INE1", "buy_sell": "BUY", "quantity": 10, "net_amount": 1000.0}]
    q = get_data_quality_summary(txs)
    assert q["imported_transactions"] == 1
    assert "missing_periods" in q

def test_xirr_audit():
    txs = [{"trade_date": "2024-01-01", "isin": "INE1", "buy_sell": "BUY", "quantity": 10, "net_amount": 1000.0}, {"trade_date": "2025-01-01", "isin": "INE1", "buy_sell": "SELL", "quantity": 10, "net_amount": 1200.0}]
    def pl(sid, d): return 120.0
    df = get_xirr_audit(txs, security_id="INE1", end_date="2025-01-01", price_lookup=pl)
    assert not df.empty
    assert {"date", "amount"}.issubset(df.columns)

def test_backup_restore(tmp_path):
    import sqlite3
    db = tmp_path / "t.db"
    conn = sqlite3.connect(str(db))
    conn.execute("CREATE TABLE t (x INTEGER)")
    conn.execute("INSERT INTO t VALUES (1)")
    conn.commit(); conn.close()
    backup = tmp_path / "b.db"
    backup_database(str(db), str(backup))
    assert backup.exists()
    # modify original
    conn = sqlite3.connect(str(db))
    conn.execute("DELETE FROM t")
    conn.commit(); conn.close()
    restore_database(str(backup), str(db))
    conn = sqlite3.connect(str(db))
    assert conn.execute("SELECT COUNT(*) FROM t").fetchone()[0] == 1
    conn.close()

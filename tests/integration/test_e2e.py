
import pandas as pd

from accounting.fifo import calculate_realized_gains, get_fifo_lots
from ledger.positions import get_holdings
from performance.benchmarks import calculate_benchmark_xirr
from performance.xirr import calculate_stock_xirr
from reports.exports import build_workbook, workbook_to_bytes
from src.holdings import load_holdings


def test_holdings_import():
    h = load_holdings("tests/fixtures/holdings.tsv")
    assert len(h) == 1
    assert h.iloc[0]["isin"] == "INEA"

def test_ledger_reconstruct():
    txs = _load_txs()
    h = get_holdings(txs, "2022-06-01")
    assert h.get("INEA", 0) == 3

def test_fifo_lots():
    txs = _load_txs()
    lots = get_fifo_lots(txs, as_of_date="2022-06-01")
    assert sum(l["qty"] for l in lots) == 3

def test_realized_gains():
    txs = _load_txs()
    gains = calculate_realized_gains(txs, start_date="2018-01-01", end_date="2022-12-31")
    assert len(gains) >= 2

def test_stock_xirr():
    txs = _load_txs()
    def pl(sid, d): return 120.0
    res = calculate_stock_xirr(txs, security_id="INEA", end_date="2022-06-01", price_lookup=pl)
    assert res["xirr"] is not None

def test_benchmark_xirr():
    def idx(d): return 1000.0 + (d.year - 2018) * 50.0 + d.timetuple().tm_yday
    cfs = [("2018-01-10", -1000.0)]
    res = calculate_benchmark_xirr(cfs, idx, "2018-01-10", "2022-06-01")
    assert res["xirr"] is not None

def test_report_opens():
    import io

    from openpyxl import load_workbook
    wb = build_workbook({"A": pd.DataFrame({"x": [1]})}, metadata={"k": "v"})
    b = workbook_to_bytes(wb)
    wb2 = load_workbook(io.BytesIO(b))
    assert "Metadata" in wb2.sheetnames

def _load_txs():
    import json
    with open("tests/fixtures/transactions.json") as f:
        return json.load(f)

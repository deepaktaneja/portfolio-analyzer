"""Report export helpers (Stage 10)."""
from __future__ import annotations

import os
import shutil
from datetime import UTC, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows


def _meta_sheet(wb, metadata):
    ws = wb.create_sheet("Metadata")
    ws.append(["Field", "Value"])
    for k, v in (metadata or {}).items():
        ws.append([str(k), str(v)])
    ws.append(["Generated at", datetime.now(UTC).isoformat()])

def build_workbook(sheets, metadata=None):
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    _meta_sheet(wb, metadata or {})
    for name, df in (sheets or {}).items():
        ws = wb.create_sheet(str(name)[:31])
        if isinstance(df, pd.DataFrame):
            for row in dataframe_to_rows(df, index=False, header=True):
                ws.append(row)
        elif isinstance(df, list):
            for row in df:
                ws.append(list(row))
        else:
            ws.append([str(df)])
    return wb

def workbook_to_bytes(wb):
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def make_metadata(start_date=None, end_date=None, methodology="", assumptions=None, sources=None, warnings=None):
    return {
        "Analysis start date": start_date,
        "Analysis end date": end_date,
        "Return methodology": methodology,
        "Assumptions": "; ".join(assumptions or []),
        "Data sources": "; ".join(sources or []),
        "Missing-data warnings": "; ".join(warnings or []),
    }

def get_data_quality_summary(transactions, holdings=None, market_db="data/market_data.sqlite3", dividends=None, actions=None):
    txs = transactions or []
    unresolved = sorted({t.get("isin") for t in txs if t.get("isin") and (not t.get("isin") or len(t.get("isin")) != 12)})
    try:
        from market_data.repository import MarketDataStore
        store = MarketDataStore(market_db)
        report = store.get_market_data_quality_report()
        missing_prices = len(report.get("anomalies", []))
    except Exception:  # noqa: BLE001
        missing_prices = 0
    unverified_divs = 0
    if dividends:
        for d in dividends:
            st = d.get("status") if isinstance(d, dict) else getattr(d, "status", None)
            if st and st != "confirmed":
                unverified_divs += 1
    pending_ca = 0
    if actions:
        for a in actions:
            if not getattr(a, "verified", True):
                pending_ca += 1
    from src.reconcile import reconcile
    recon_diffs = 0
    if holdings is not None:
        try:
            recon_diffs = len(reconcile(txs, holdings))
        except Exception:  # noqa: BLE001
            recon_diffs = 0
    from src.pdf_parser import detect_duplicate_details, detect_missing_quarters
    missing_periods = detect_missing_quarters(txs)
    dup = detect_duplicate_details(txs)
    return {
        "imported_transactions": len(txs),
        "unresolved_identifiers": unresolved,
        "missing_price_anomalies": missing_prices,
        "unverified_dividends": unverified_divs,
        "pending_corporate_actions": pending_ca,
        "holdings_reconciliation_differences": recon_diffs,
        "missing_periods": missing_periods,
        "duplicate_or_conflicting_records": len(dup),
    }

def get_xirr_audit(transactions, actions=None, dividends=None, price_lookup=None, security_id=None, start_date=None, end_date=None, scope="stock"):
    from performance.xirr import get_xirr_cashflow_audit
    cfs = get_xirr_cashflow_audit(transactions, actions, dividends, price_lookup, security_id=security_id, start_date=start_date, end_date=end_date, scope=scope)
    rows = []
    for c in cfs if isinstance(cfs, list) else [cfs]:
        if isinstance(c, (list, tuple)) and len(c) >= 2:
            rows.append({"date": c[0], "amount": c[1], "label": c[2] if len(c) > 2 else ""})
    return pd.DataFrame(rows)

def backup_database(db_path, dest):
    shutil.copy2(db_path, dest)
    return dest

def restore_database(backup, db_path):
    if not os.path.exists(backup):
        raise FileNotFoundError(backup)
    shutil.copy2(backup, db_path)
    return db_path


def make_report(sheets, metadata=None):
    wb = build_workbook(sheets, metadata)
    return workbook_to_bytes(wb)
